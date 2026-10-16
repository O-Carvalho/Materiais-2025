import pandas as pd, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from familias import familia

SAIDA = sys.argv[1]
ANO = os.environ.get("ANO", "2025")
ESTOQUE = "Não gerou necessidade cotação"
MOT_EST = "Havia estoque, mas o item não foi baixado pelo almoxarifado"
MOT_COMP = "Parado em compras (Cotar / Em Cotação / Pedido)"
c = pd.read_excel("Consolidado OS.xlsx")
d = c.sort_values(["OS", "Material"]).reset_index(drop=True)
d["Família"] = d.Material.map(familia)
d["_gr"] = d["Situação"].map({"Baixada": "Entregue", "Cancelada": "Cancelado"}).fillna("Em aberto")
N = len(d); U = N + 1

FONTE = "Arial"; AZUL, TINTA, CINZA, LINHA, VERM = "2A78D6", "0B0B0B", "52514E", "D9D8D4", "C00000"
f_tit = Font(name=FONTE, size=18, bold=True, color=TINTA)
f_sec = Font(name=FONTE, size=12, bold=True, color=TINTA)
f_sub = Font(name=FONTE, size=10, color=CINZA)
f_hdr = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
f_txt = Font(name=FONTE, size=10, color=TINTA)
f_bold = Font(name=FONTE, size=10, bold=True, color=TINTA)
fill_hdr = PatternFill("solid", fgColor="1C5CAB")
NUM, QTD, PCT = "#,##0", "#,##0.##", "0.0%"

wb = Workbook(); dash = wb.active; dash.title = "Dashboard"

def cabecalho(ws, cols, larg, row=1):
    for i, (h, w) in enumerate(zip(cols, larg), 1):
        cel = ws.cell(row, i, h); cel.font = f_hdr; cel.fill = fill_hdr
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[row].height = 30; ws.freeze_panes = ws.cell(row + 1, 1)

def tabela(ws, nome, ncol, r2, r1=1):
    t = Table(displayName=nome, ref=f"A{r1}:{L(ncol)}{r2}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); ws.add_table(t)

def fonte(ws, r1, r2, c2):
    for row in ws.iter_rows(min_row=r1, max_row=r2, max_col=c2):
        for cel in row:
            if cel.font != f_bold: cel.font = f_txt

# ---------- Base do ano (base extraída, mantida + 3 colunas de apoio) ----------
B = wb.create_sheet(f"Base {ANO}")
orig = list(c.columns)
cols = orig + ["Família", "Classificação", "Motivo (em aberto)"]
cabecalho(B, cols, [8,11,12,12,62,11,12,11,28,17,11,11,11,30,13,52])
for i, row in enumerate(d[orig + ["Família"]].itertuples(index=False), 2):
    for j, v in enumerate(row, 1):
        if pd.isna(v): continue
        B.cell(i, j, v.to_pydatetime() if j == 10 else v)
    B.cell(i, 15, f'=IF(M{i}="Baixada","Entregue",IF(M{i}="Cancelada","Cancelado","Em aberto"))')
    B.cell(i, 16, f'=IF(O{i}<>"Em aberto","",IF(I{i}="{ESTOQUE}","{MOT_EST}","{MOT_COMP}"))')
    B.cell(i, 10).number_format = "dd/mm/yyyy hh:mm"
    for j in (6, 7, 8): B.cell(i, j).number_format = QTD
fonte(B, 2, U, 16); tabela(B, f"Base{ANO}", 16, U)
B["J1"].comment = Comment("Data em que o item foi solicitado no sistema. Não é usada nas análises.", "Relatório")
R = lambda col: f"'Base {ANO}'!${col}$2:${col}${U}"
OSr, QTDr, BXr, FAMr, CLr, MOTr = R("A"), R("F"), R("G"), R("N"), R("O"), R("P")

# ---------- Ranking de materiais ----------
P = wb.create_sheet("Ranking Materiais")
e = d[d._gr != "Cancelado"]
g = e.groupby("Cod. Material").agg(Material=("Material","first"), Fam=("Família","first"), nos=("OS","nunique"),
        req=("OS","size"), sol=("Quantidade","sum"), bx=("Quantidade Baixada","sum"))
g = g.sort_values(["nos","bx"], ascending=False).reset_index()
cabecalho(P, ["Posição","Cod. Material","Material","Família","Nº de OS que usaram","Qtd. solicitada","Qtd. baixada (consumida)"], [8,12,62,30,12,13,14])
for i, r in enumerate(g.itertuples(index=False), 2):
    vals = [i-1, int(r[0]), r.Material, r.Fam, int(r.nos), float(r.sol), float(r.bx)]
    for j, v in enumerate(vals, 1): P.cell(i, j, v)
    P.cell(i,5).number_format = NUM; P.cell(i,6).number_format = QTD; P.cell(i,7).number_format = QTD
np_ = len(g) + 1
fonte(P, 2, np_, 7); tabela(P, "Ranking", 7, np_)
P["E1"].comment = Comment("Ordenado pelo nº de OS que usaram o material (sem contar cancelados). Desempate pela quantidade baixada.", "Relatório")
P["G1"].comment = Comment("Na unidade do próprio material (peça, metro, kg...). Não compare entre materiais de unidades diferentes.", "Relatório")
P.cell(1,9,"Material (resumo)").font = f_hdr; P.cell(1,9).fill = fill_hdr; P.column_dimensions["I"].width = 40
for i in range(2, 12): P.cell(i, 9, f"=LEFT(C{i},38)").font = f_txt

# ---------- Itens em Aberto ----------
A = wb.create_sheet("Itens em Aberto")
ab = d[d._gr == "Em aberto"].copy()
ab["Motivo"] = ab["Status Necessidade"].map(lambda s: MOT_EST if s == ESTOQUE else MOT_COMP)
ab["Obs"] = ab.apply(lambda r: f"Baixa parcial: {r['Quantidade Baixada']:g} de {r['Quantidade']:g}" if r["Quantidade Baixada"] > 0 else "", axis=1)
ab = ab.sort_values(["Motivo", "OS"])
cols_a = ["OS","Cod. Material","Material","Família","Status Necessidade","Quantidade","Quantidade Baixada","Motivo","Obs"]
cabecalho(A, cols_a, [8,12,62,28,26,11,12,50,26])
for i, row in enumerate(ab[cols_a].itertuples(index=False), 2):
    for j, v in enumerate(row, 1): A.cell(i, j, v)
    A.cell(i,6).number_format = QTD; A.cell(i,7).number_format = QTD
na = len(ab) + 1
fonte(A, 2, na, 9); tabela(A, "EmAberto", 9, na)

# ---------- Por OS ----------
O = wb.create_sheet("Por OS")
oss = sorted(d.OS.unique())
cabecalho(O, ["OS","Itens requisitados","Itens entregues","Itens cancelados","Itens em aberto","Qtd. baixada","% entregue","Situação da OS"], [8,13,13,13,13,14,11,16])
for i, n in enumerate(oss, 2):
    O.cell(i,1,int(n))
    O.cell(i,2,f"=COUNTIFS({OSr},A{i})")
    O.cell(i,3,f'=COUNTIFS({OSr},A{i},{CLr},"Entregue")')
    O.cell(i,4,f'=COUNTIFS({OSr},A{i},{CLr},"Cancelado")')
    O.cell(i,5,f'=COUNTIFS({OSr},A{i},{CLr},"Em aberto")')
    O.cell(i,6,f'=SUMIFS({BXr},{OSr},A{i},{CLr},"<>Cancelado")').number_format = QTD
    O.cell(i,7,f"=IFERROR(C{i}/(B{i}-D{i}),0)").number_format = PCT
    O.cell(i,8,f'=IF(E{i}>0,"Com pendência","OK")')
no = len(oss) + 1
fonte(O, 2, no, 8); tabela(O, "PorOS", 8, no)
O["G1"].comment = Comment("Itens entregues ÷ (itens requisitados − cancelados).", "Relatório")

# ---------- Por Família ----------
F = wb.create_sheet("Por Família")
fam = d[d._gr == "Entregue"].groupby("Família").size().sort_values(ascending=False).index.tolist()
fam += [x for x in d["Família"].unique() if x not in fam]
cabecalho(F, ["Família","Itens requisitados","Itens entregues","Itens cancelados","Itens em aberto","Qtd. baixada","% dos itens entregues"], [34,13,13,13,13,14,13])
for i, f_ in enumerate(fam, 2):
    F.cell(i,1,f_)
    F.cell(i,2,f"=COUNTIFS({FAMr},A{i})")
    F.cell(i,3,f'=COUNTIFS({FAMr},A{i},{CLr},"Entregue")')
    F.cell(i,4,f'=COUNTIFS({FAMr},A{i},{CLr},"Cancelado")')
    F.cell(i,5,f'=COUNTIFS({FAMr},A{i},{CLr},"Em aberto")')
    F.cell(i,6,f'=SUMIFS({BXr},{FAMr},A{i},{CLr},"<>Cancelado")')
    F.cell(i,7,f"=IFERROR(C{i}/SUM($C$2:$C${len(fam)+1}),0)")
    for j, fmt in zip(range(2,8), [NUM,NUM,NUM,NUM,QTD,PCT]): F.cell(i,j).number_format = fmt
nf = len(fam) + 1; t = nf + 1
F.cell(t,1,"Total").font = f_bold
for j in range(2,8):
    F.cell(t,j,f"=SUM({L(j)}2:{L(j)}{nf})").font = f_bold; F.cell(t,j).number_format = [NUM,NUM,NUM,NUM,QTD,PCT][j-2]
fonte(F, 2, nf, 7); tabela(F, "PorFamilia", 7, nf)
F.cell(t+2,1,"A Qtd. baixada soma unidades diferentes (peças, metros, kg). Serve para comparar a mesma família ao longo do tempo, não famílias entre si.").font = f_sub

# ---------- Critérios ----------
C = wb.create_sheet("Critérios")
C.column_dimensions["A"].width = 26; C.column_dimensions["B"].width = 110
nbug = int(((d["Situação"]=="Baixada") & (d["Status Necessidade"]!=ESTOQUE)).sum())
linhas = [
 ("Critérios do relatório", None),
 ("Escopo", f"Todas as {d.OS.nunique()} OS de {ANO} extraídas do ERP, com todos os seus itens. A coluna Data (data de solicitação no sistema) não é usada como filtro nem em nenhuma análise."),
 ("Entregue", "Situação = Baixada. O almoxarifado baixou o item e entregou à produção."),
 ("Cancelado", "Situação = Cancelada. O material deixou de ser necessário. Fica fora do consumo."),
 ("Em aberto", "Situação = Aberto. Indica erro no processo. O motivo vem do Status Necessidade:"),
 ("", f"• Não gerou necessidade cotação → {MOT_EST}."),
 ("", f"• Cotar / Em Cotação / Pedido → {MOT_COMP}."),
 ("Bug do ERP", f"{nbug} itens importados com Status = Pedido e Situação = Baixada (erro do sistema). Contam como entregues, não como problema."),
 ("Consumo (Qtd. baixada)", "Soma da Quantidade Baixada dos itens não cancelados, incluindo baixas parciais de itens em aberto."),
 ("Ranking", "Materiais ordenados pelo nº de OS que usaram o material. Desempate pela quantidade baixada."),
 ("Família", f"Agrupamento sugerido pela descrição do material (ex.: TERMINAL, BORNE → Conexão e bornes). Ajustável na coluna Família da Base {ANO}."),
 ("Unidades", "O ERP não informa a unidade de medida. Compare quantidades só dentro do mesmo material."),
]
for i, (a, b) in enumerate(linhas, 1):
    C.cell(i,1,a).font = Font(name=FONTE, bold=True, size=13 if i == 1 else 10); C.cell(i,1).alignment = Alignment(vertical="top")
    if b: C.cell(i,2,b).font = f_txt; C.cell(i,2).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- Dashboard ----------
dash.sheet_view.showGridLines = False
for col in range(1, 30): dash.column_dimensions[L(col)].width = 9.5
dash.column_dimensions["A"].width = 2
dash["B1"] = f"Consumo de Materiais – {ANO}"; dash["B1"].font = f_tit
dash["B2"] = f"Todas as {d.OS.nunique()} OS de {ANO} extraídas do ERP · {N:,} itens".replace(",", "."); dash["B2"].font = f_sub
borda = Side(style="thin", color=LINHA)
kpis = [
 ("OS atendidas", f"=COUNT('Por OS'!A2:A{no})", "ordens de serviço", TINTA),
 ("Itens entregues", f'=COUNTIF({CLr},"Entregue")', f'="de "&TEXT(COUNTA({OSr}),"#,##0")&" requisitados"', TINTA),
 ("Itens em aberto", f'=COUNTIF({CLr},"Em aberto")', "erros no processo", VERM),
 ("OS com pendência", f"=COUNTIF('Por OS'!H2:H{no},\"Com pendência\")", "têm item em aberto", VERM),
]
for k, (rot, form, nota, cor) in enumerate(kpis):
    c0 = 2 + k * 6
    for r in (4, 5, 6):
        dash.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 4)
        for cc in range(c0, c0 + 5):
            dash.cell(r, cc).fill = PatternFill("solid", fgColor="F4F3F0")
            dash.cell(r, cc).border = Border(top=borda if r == 4 else None, bottom=borda if r == 6 else None,
                                             left=borda if cc == c0 else None, right=borda if cc == c0 + 4 else None)
        dash.cell(r, c0).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    dash.cell(4, c0, rot).font = Font(name=FONTE, size=10, bold=True, color=CINZA)
    v = dash.cell(5, c0, form); v.number_format = NUM; v.font = Font(name=FONTE, size=22, bold=True, color=cor)
    dash.cell(6, c0, nota).font = Font(name=FONTE, size=9, color=CINZA)
dash.row_dimensions[5].height = 34

def barra(titulo, ws, cat_col, val_col, r1, r2, alt=11, larg=21.5):
    ch = BarChart(); ch.type = "bar"; ch.title = titulo; ch.style = 2; ch.legend = None
    ch.add_data(Reference(ws, min_col=val_col, min_row=r1, max_row=r2), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=r1, max_row=r2))
    s = ch.series[0]; s.graphicalProperties = GraphicalProperties(solidFill=AZUL); s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList(); s.dLbls.showVal = True
    for a in ("showSerName", "showCatName", "showLegendKey", "showPercent"): setattr(s.dLbls, a, False)
    ch.gapWidth = 60; ch.y_axis.majorGridlines = None; ch.y_axis.delete = True
    ch.x_axis.delete = False; ch.x_axis.tickLblSkip = 1; ch.x_axis.scaling.orientation = "maxMin"
    ch.height = alt; ch.width = larg
    return ch

dash.add_chart(barra("Top 10 materiais mais consumidos (nº de OS)", P, 9, 5, 2, 11), "B8")
dash.add_chart(barra("Itens entregues por família", F, 1, 3, 2, nf), "N8")

# Quadro de itens em aberto por motivo
r0 = 31
dash.cell(r0, 2, "Itens em aberto por motivo").font = f_sec
for j, h in zip((2, 10, 11, 12), ("Motivo", "Itens", "", "OS")):
    cel = dash.cell(r0 + 1, j, h); cel.font = f_hdr; cel.fill = fill_hdr
dash.merge_cells(start_row=r0+1, start_column=2, end_row=r0+1, end_column=9)
for k, m in enumerate((MOT_EST, MOT_COMP)):
    r = r0 + 2 + k
    dash.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
    dash.cell(r, 2, m).font = f_txt
    dash.cell(r, 10, f'=COUNTIF({MOTr},B{r})').font = f_bold
    dash.cell(r, 12, int(ab[ab.Motivo == m].OS.nunique())).font = f_bold
    for cc in range(2, 13): dash.cell(r, cc).border = Border(bottom=borda)
dash.cell(r0 + 4, 2, "Lista completa, com OS e material, na aba 'Itens em Aberto'.").font = f_sub
dash.cell(r0 + 1, 12).comment = Comment("Nº de OS distintas com itens nesse motivo (calculado na geração).", "Relatório")

ordem = ["Dashboard", "Ranking Materiais", "Itens em Aberto", "Por OS", "Por Família", f"Base {ANO}", "Critérios"]
wb._sheets = [wb[n] for n in ordem]; wb.active = 0
if os.environ.get("PREVIEW"):
    for ws in wb.worksheets:
        if ws.title != "Dashboard": ws.sheet_state = "hidden"
    dash.print_area = "A1:Y37"; dash.page_setup.orientation = "landscape"
    dash.sheet_properties.pageSetUpPr.fitToPage = True; dash.page_setup.fitToWidth = 1; dash.page_setup.fitToHeight = 1
wb.save(SAIDA)
print("ok", N, "linhas,", len(oss), "OS,", len(ab), "em aberto,", nbug, "bug")
