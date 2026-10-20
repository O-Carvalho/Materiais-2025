"""Organiza a base de materiais baixados dos painéis em montagem (uma aba, ordenada por OS)."""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter as L

COLS = ["OS", "Documento", "Nº Solicitação", "Cod. Material", "Material", "Quantidade", "Quantidade Baixada",
        "Necessidade", "Status Necessidade", "Data", "Cód. Aplicação", "Cód. Localização", "Situação"]
LARG = [8, 11, 12, 12, 70, 11, 12, 11, 28, 17, 11, 11, 11]

def gera(entradas, saida):
    d = pd.concat([pd.read_excel(p) for p in entradas], ignore_index=True)
    falta = [c for c in COLS if c not in d.columns]
    if falta: raise SystemExit(f"Colunas ausentes: {falta}")
    d = d[COLS].drop_duplicates().sort_values(["OS", "Documento", "Material"]).reset_index(drop=True)
    wb = Workbook(); ws = wb.active; ws.title = "Base"
    F = "Arial"
    for j, (h, w) in enumerate(zip(COLS, LARG), 1):
        x = ws.cell(1, j, h); x.font = Font(name=F, bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", fgColor="1C5CAB"); x.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[1].height = 30; ws.freeze_panes = "A2"
    for i, row in enumerate(d.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if pd.isna(v): continue
            if j == 10: v = v.to_pydatetime()
            elif isinstance(v, float) and v.is_integer(): v = int(v)
            ws.cell(i, j, v).font = Font(name=F, size=10)
        ws.cell(i, 10).number_format = "dd/mm/yyyy"
        for j in (6, 7): ws.cell(i, j).number_format = "#,##0.##"
    t = Table(displayName="BaseMontagem", ref=f"A1:{L(len(COLS))}{len(d)+1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); ws.add_table(t)
    wb.save(saida)
    return d

if __name__ == "__main__":
    d = gera(sys.argv[2:], sys.argv[1])
    print(len(d), "linhas,", d.OS.nunique(), "OS,", d["Cod. Material"].nunique(), "materiais")
