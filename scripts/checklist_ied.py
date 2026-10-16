"""Gera a planilha de conferência de IED: uma linha por OS com placa frontal com recorte para IED."""
import re, sys, warnings
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from ied import analisa
warnings.filterwarnings("ignore")

PROJ = r"(SPCS[_ ][0-9_]+\d|PSI\.\w+[_ ][0-9_]+\d|PCS[_ ][0-9_]+\d)"

def projeto_tag(desc):
    d = re.sub(r"^SEMI ACAB\w*\s+", "", str(desc).strip(), flags=re.I)
    p = re.search(PROJ, d, re.I)
    proj = p.group(1).strip("_ ") if p else ""
    partes = [x.strip() for x in re.split(r"\s[-–]\s?|\s?[-–]\s", d) if x.strip()]
    ult = partes[-1] if len(partes) > 1 else d
    m = re.search(r"\b(?:PN|PNL|PAINEL)\s+(.+)$", ult, re.I)
    if m:
        tag = m.group(1)
    elif len(partes) > 1 and not re.search(r"\b(SE|ETD|CTEEP)\b", ult):
        tag = ult
    else:
        tag = ult.split()[-1]
    return proj, tag.strip()

def gera(saida):
    c = pd.read_excel("Consolidado OS.xlsx")
    r, _ = analisa(c)
    r = r[r.Vagas_IED > 0].copy()
    r["Tem IED"] = ((r.SIPROTEC + r.Outros_IED) > 0).map({True: "Sim", False: "Não"})
    r[["Projeto", "TAG"]] = r.Semiacabado.apply(lambda s: pd.Series(projeto_tag(s) if s else ("", "")))
    r = r.sort_values(["Tem IED", "Projeto", "OS"])  # "Não" primeiro
    wb = Workbook(); ws = wb.active; ws.title = "Conferência IED"
    cols = ["OS", "Descrição", "Projeto", "TAG", "Tem placa p/ IED?", "Tem IED?", "Faz sentido?", "Observação"]
    larg = [8, 72, 20, 14, 11, 10, 18, 45]
    F = "Arial"; hdr = PatternFill("solid", fgColor="1C5CAB"); pend = PatternFill("solid", fgColor="FFF4CC")
    for j, (h, w) in enumerate(zip(cols, larg), 1):
        x = ws.cell(1, j, h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = hdr
        x.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[chr(64 + j)].width = w
    ws.row_dimensions[1].height = 30; ws.freeze_panes = "B2"
    for i, row in enumerate(r.itertuples(index=False), 2):
        tem = r["Tem IED"].iloc[i - 2]
        vals = [row.OS, row.Semiacabado or "(OS sem semiacabado)", row.Projeto, row.TAG, "Sim", tem,
                "Faz sentido" if tem == "Sim" else None, None]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v).font = Font(name=F, size=10)
        if vals[5] == "Não": ws.cell(i, 7).fill = pend
    n = len(r) + 1
    dv = DataValidation(type="list", formula1='"Faz sentido,Não faz sentido"', allow_blank=True)
    dv.prompt = "Escolha: Faz sentido / Não faz sentido"; dv.promptTitle = "Faz sentido?"
    ws.add_data_validation(dv); dv.add(f"G2:G{n}")
    t = Table(displayName="ConferenciaIED", ref=f"A1:H{n}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); ws.add_table(t)
    wb.save(saida)
    return r

if __name__ == "__main__":
    r = gera(sys.argv[1])
    print(len(r), "OS com placa;", (r["Tem IED"] == "Sim").sum(), "com IED;", (r["Tem IED"] == "Não").sum(), "sem IED")
