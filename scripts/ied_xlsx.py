import pandas as pd, sys, warnings
warnings.filterwarnings("ignore")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter as L
from ied import analisa

SAIDA = sys.argv[1]
c = pd.read_excel("Consolidado OS.xlsx")
r, cc = analisa(c)
ORDEM = ["ALERTA – painel preparado p/ IED, sem IED", "IED a menos que o previsto", "Verificar – painel sem IED no sistema",
         "IED cancelado", "OK – tem IED", "Sem IED – ENEL (IED de terceiro)",
         "Sem IED – painel auxiliar (IHM, SA, telecom, DGO)", "Sem IED – conjunto TC/TP (não leva IED)"]
EXPL = {
 ORDEM[0]: "Tem placa frontal (PFF) com recorte para IED, mas nenhum IED foi requisitado no sistema. É o caso mais estranho.",
 ORDEM[1]: "Tem IED, mas menos do que os recortes de IED nas placas frontais.",
 ORDEM[2]: "Sem IED e sem placa com recorte para IED. Pelo nome não dá para saber se o painel leva IED: conferir no DB Projetos.",
 ORDEM[3]: "O IED foi requisitado e depois cancelado.",
 ORDEM[4]: "IED requisitado em quantidade igual ou maior que os recortes.",
 ORDEM[5]: "Painel ENEL: o IED vem de terceiro e não entra no sistema.",
 ORDEM[6]: "IHM, serviço auxiliar (QDCA/QDCC), telecom, DGO, caixa de interligação: normalmente não levam IED.",
 ORDEM[7]: "Projetos PSI.EPC (caixas e conjuntos de TC/TP): não levam IED.",
}
r["_o"] = r.Diagnostico.map({k: i for i, k in enumerate(ORDEM)})
r = r.sort_values(["_o", "Cliente", "OS"])

F = "Arial"
f_tit = Font(name=F, size=16, bold=True); f_sec = Font(name=F, size=12, bold=True); f_sub = Font(name=F, size=10, color="52514E")
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF"); f_txt = Font(name=F, size=10); f_b = Font(name=F, size=10, bold=True)
fill = PatternFill("solid", fgColor="1C5CAB")
CORES = {0: "F8D7D7", 1: "FCE8D5", 2: "FFF4CC"}

wb = Workbook(); R = wb.active; R.title = "Resumo"

def cab(ws, row, cols, larg=None):
    for j, h in enumerate(cols, 1):
        x = ws.cell(row, j, h); x.font = f_hdr; x.fill = fill; x.alignment = Alignment(wrap_text=True, vertical="center")
        if larg: ws.column_dimensions[L(j)].width = larg[j-1]
    ws.row_dimensions[row].height = 30

# ---- Por OS ----
P = wb.create_sheet("Por OS")
cols = ["OS","Projeto","Cliente","Painel (semiacabado)","Recortes p/ IED (PFF)","SIPROTEC requisitados","SIPROTEC baixados","SIPROTEC cancelados","Outros IED (SEL/Hitachi)","Modelos SIPROTEC","Diagnóstico"]
keys = ["OS","Projeto","Cliente","Semiacabado","Vagas_IED","SIPROTEC","SIPROTEC_baixado","SIPROTEC_cancelado","Outros_IED","Modelos","Diagnostico"]
cab(P, 1, cols, [8,11,16,70,11,11,11,11,12,22,46]); P.freeze_panes = "B2"
for i, row in enumerate(r[keys].itertuples(index=False), 2):
    for j, v in enumerate(row, 1):
        x = P.cell(i, j, (int(v) if isinstance(v, float) and v.is_integer() else v) if v != "" else None); x.font = f_txt
    o = ORDEM.index(row[-1])
    if o in CORES: P.cell(i, 11).fill = PatternFill("solid", fgColor=CORES[o])
n = len(r) + 1
t = Table(displayName="IEDporOS", ref=f"A1:{L(len(cols))}{n}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); P.add_table(t)

# ---- Itens SIPROTEC ----
S = wb.create_sheet("Itens SIPROTEC")
it = cc[cc._sip | cc._outro].sort_values(["OS"])
ci = ["OS","Cod. Material","Material","Quantidade","Quantidade Baixada","Status Necessidade","Situação"]
cab(S, 1, ci, [8,12,70,11,12,28,11]); S.freeze_panes = "A2"
for i, row in enumerate(it[ci].itertuples(index=False), 2):
    for j, v in enumerate(row, 1): S.cell(i, j, v).font = f_txt
t = Table(displayName="ItensIED", ref=f"A1:G{len(it)+1}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); S.add_table(t)

# ---- Resumo ----
R.sheet_view.showGridLines = False
R.column_dimensions["A"].width = 2; R.column_dimensions["B"].width = 46; R.column_dimensions["C"].width = 8; R.column_dimensions["D"].width = 95
R["B1"] = "Verificação de IEDs SIPROTEC por OS"; R["B1"].font = f_tit
R["B2"] = f"{r.OS.nunique()} OS da base · {int(r.SIPROTEC.sum())} SIPROTEC requisitados (não cancelados) em {int((r.SIPROTEC>0).sum())} OS · {int(r.Vagas_IED.sum())} recortes para IED nas placas frontais"
R["B2"].font = f_sub
R["B4"] = "OS por diagnóstico"; R["B4"].font = f_sec
for j, h in zip((2, 3, 4), ("Diagnóstico", "OS", "O que significa")):
    x = R.cell(5, j, h); x.font = f_hdr; x.fill = fill
pos = "'Por OS'!$K$2:$K$" + str(n)
for k, d in enumerate(ORDEM):
    i = 6 + k
    R.cell(i, 2, d).font = f_b if k < 3 else f_txt
    R.cell(i, 3, f'=COUNTIF({pos},B{i})').font = f_b
    R.cell(i, 4, EXPL[d]).font = f_txt; R.cell(i, 4).alignment = Alignment(wrap_text=True, vertical="top")
    if k in CORES: R.cell(i, 2).fill = PatternFill("solid", fgColor=CORES[k])
R.cell(14, 2, "Total").font = f_b; R.cell(14, 3, "=SUM(C6:C13)").font = f_b

# por cliente (casos a verificar)
R["B16"] = "Casos a verificar (3 primeiros diagnósticos) por cliente"; R["B16"].font = f_sec
for j, h in zip((2, 3, 4), ("Cliente", "OS", "Projetos / painéis")):
    x = R.cell(17, j, h); x.font = f_hdr; x.fill = fill
v = r[r._o < 3]
g = v.groupby("Cliente").agg(os=("OS", "nunique"), proj=("Semiacabado", lambda s: s.str.upper().str.extract(r"((?:SPCS|PSI\.SIS|PSI\.EPC|PCS)[_ ][0-9_]+)")[0].dropna().str.strip("_ ").unique())).sort_values("os", ascending=False)
i = 18
for cli, row in g.iterrows():
    R.cell(i, 2, cli).font = f_txt; R.cell(i, 3, int(row.os)).font = f_b
    R.cell(i, 4, ", ".join(row.proj) if len(row.proj) else "—").font = f_txt; i += 1

i += 1
R.cell(i, 2, "Como foi feito").font = f_sec
notas = [
 "SIPROTEC: itens 'RELE DIGITAL / RELE DE PROTECAO' da Siemens com código 7SA, 7SJ, 7SL, 7SD, 7SS, 7UT, 7VK… (o ERP não usa a palavra SIPROTEC).",
 "Recortes p/ IED: placas frontais PFF com '1 IED' ou '2 IED' na descrição (ex.: PFF 19\" 07U 1 IED SIP5 UNIVERSAL). Mostram que o painel foi projetado para receber IED.",
 "Projeto, cliente e painel: lidos do item SEMI ACABADO de cada OS (ex.: SPCS_201_23_083 - CTEEP - SE CACONDE PAINEL PPTR3).",
 "Itens cancelados não contam como IED requisitado.",
 "Limite: sem o campo 'tipo de painel' do DB Projetos não dá para separar com certeza os painéis que levam IED dos que não levam, nem os painéis de não energia.",
]
for k, t_ in enumerate(notas, i + 1):
    R.cell(k, 2, "• " + t_).font = f_txt
wb._sheets = [R, P, S]; wb.active = 0
wb.save(SAIDA)
print(r.Diagnostico.value_counts().reindex(ORDEM).to_string())
print(g.to_string())
