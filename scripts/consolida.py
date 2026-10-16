import glob, re, xlrd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

arquivos = sorted(glob.glob("OS *.xls"), key=lambda f: int(re.search(r"\d+", f).group()))
wb = Workbook(); ws = wb.active; ws.title = "Consolidado"
cab = None; n = 0
for f in arquivos:
    b = xlrd.open_workbook(f); s = b.sheet_by_index(0)
    h = s.row_values(0)
    if cab is None:
        cab = h; ws.append(h); dc = cab.index("Data")
    assert h == cab, f
    for r in range(1, s.nrows):
        linha = []
        for c in range(s.ncols):
            t, v = s.cell_type(r, c), s.cell_value(r, c)
            if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or v == "":
                v = None
            elif t == xlrd.XL_CELL_DATE:
                v = xlrd.xldate_as_datetime(v, b.datemode)
            elif t == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                v = int(v)
            linha.append(v)
        ws.append(linha); n += 1
        ws.cell(n + 1, dc + 1).number_format = "dd/mm/yyyy hh:mm:ss"
larg = {"OS": 9, "Material": 70, "Status Necessidade": 32, "Data": 20}
for i, nome in enumerate(cab, 1):
    ws.column_dimensions[get_column_letter(i)].width = larg.get(nome, max(14, len(nome) + 4))
t = Table(displayName="BaseOS", ref=f"A1:{get_column_letter(len(cab))}{n + 1}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(t); ws.freeze_panes = "A2"
wb.save("Consolidado OS.xlsx")
print(len(arquivos), "arquivos,", n, "linhas")
